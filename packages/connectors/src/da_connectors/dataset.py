"""数据集接入（3.2 快速接入的零门槛档）：上传 CSV/TSV/Excel → SQLite 表 → 即可提问。

Julius 式冷启动：第一分钟传文件就能用，产生价值后再引导连数仓（12 章 TTFV 北极星）。
文件落成 SQLite 表后复用 SQLiteConnector 全套能力（护栏/权限/profiling/冷启动）。
类型推断：整列可解析为 int→INTEGER，float→REAL，否则 TEXT。
"""

from __future__ import annotations

import csv
import io
import re
import sqlite3
from dataclasses import dataclass
from pathlib import Path

from da_connectors.sqlite import SQLiteConnector

MAX_COLUMNS = 128


@dataclass
class IngestResult:
    table: str
    columns: list[str]
    rows: int


class DatasetStore:
    """一个租户的数据集库：多个上传文件 → 同一 SQLite 文件的多张表。"""

    def __init__(self, db_path: str | Path) -> None:
        self._db_path = Path(db_path)
        self._db_path.parent.mkdir(parents=True, exist_ok=True)
        # 确保库文件存在（SQLiteConnector 以只读打开）
        sqlite3.connect(self._db_path).close()

    def connector(self, source_id: str = "datasets") -> SQLiteConnector:
        return SQLiteConnector(source_id, self._db_path)

    def list_tables(self) -> list[str]:
        conn = sqlite3.connect(self._db_path)
        try:
            return [
                r[0] for r in conn.execute(
                    "SELECT name FROM sqlite_master WHERE type='table' "
                    "AND name NOT LIKE 'sqlite_%' ORDER BY name"
                )
            ]
        finally:
            conn.close()

    def ingest_csv(
        self, content: bytes | str, table: str, delimiter: str | None = None
    ) -> IngestResult:
        text = content.decode("utf-8-sig") if isinstance(content, bytes) else content
        if delimiter is None:
            delimiter = "\t" if "\t" in text.splitlines()[0] else ","
        reader = csv.reader(io.StringIO(text), delimiter=delimiter)
        rows = [r for r in reader if any(cell.strip() for cell in r)]
        if len(rows) < 2:
            raise ValueError("数据集为空或只有表头")
        header, data = rows[0], rows[1:]
        return self._write_table(table, header, data)

    def ingest_excel(self, content: bytes, table: str, sheet: str | None = None) -> IngestResult:
        try:
            import openpyxl  # 延迟导入：optional extra "excel"
        except ImportError as e:
            raise ValueError("Excel 支持需要安装 openpyxl（uv sync --extra excel）") from e
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb[sheet] if sheet else wb.active
        rows = [
            ["" if c is None else str(c) for c in row]
            for row in ws.iter_rows(values_only=True)
        ]
        rows = [r for r in rows if any(cell.strip() for cell in r)]
        if len(rows) < 2:
            raise ValueError("数据集为空或只有表头")
        return self._write_table(table, rows[0], rows[1:])

    def _write_table(
        self, table: str, header: list[str], data: list[list[str]]
    ) -> IngestResult:
        table = _safe_identifier(table)
        columns = _dedupe([_safe_identifier(h) or f"col_{i}" for i, h in enumerate(header)])
        if len(columns) > MAX_COLUMNS:
            raise ValueError(f"列数超限（{len(columns)} > {MAX_COLUMNS}）")

        types = [_infer_type([row[i] if i < len(row) else "" for row in data])
                 for i in range(len(columns))]
        conn = sqlite3.connect(self._db_path)
        try:
            conn.execute(f'DROP TABLE IF EXISTS "{table}"')
            cols_ddl = ", ".join(f'"{c}" {t}' for c, t in zip(columns, types, strict=True))
            conn.execute(f'CREATE TABLE "{table}" ({cols_ddl})')
            placeholders = ",".join("?" * len(columns))
            normalized = [
                [_coerce(row[i] if i < len(row) else "", types[i])
                 for i in range(len(columns))]
                for row in data
            ]
            conn.executemany(
                f'INSERT INTO "{table}" VALUES ({placeholders})', normalized
            )
            conn.commit()
        finally:
            conn.close()
        return IngestResult(table=table, columns=columns, rows=len(data))


def _safe_identifier(name: str) -> str:
    cleaned = re.sub(r"[^\w一-鿿]", "_", name.strip()).strip("_")
    return cleaned[:64] or "t"


def _dedupe(names: list[str]) -> list[str]:
    seen: dict[str, int] = {}
    out = []
    for n in names:
        if n in seen:
            seen[n] += 1
            out.append(f"{n}_{seen[n]}")
        else:
            seen[n] = 0
            out.append(n)
    return out


def _infer_type(values: list[str]) -> str:
    non_empty = [v.strip() for v in values if v.strip()]
    if not non_empty:
        return "TEXT"
    if all(re.fullmatch(r"-?\d+", v) for v in non_empty):
        return "INTEGER"
    if all(re.fullmatch(r"-?\d+(\.\d+)?", v) for v in non_empty):
        return "REAL"
    return "TEXT"


def _coerce(value: str, sql_type: str):
    v = value.strip()
    if not v:
        return None
    if sql_type == "INTEGER":
        return int(v)
    if sql_type == "REAL":
        return float(v)
    return v
